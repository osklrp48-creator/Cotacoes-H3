import io
from types import SimpleNamespace

import pytest
from docx import Document
from openpyxl import Workbook

from app.models import UsoIA
from app.services.ia import leitor

RESPOSTA_IA = {
    "fornecedor": "Distribuidora Alfa",
    "cnpj": "11.222.333/0001-44",
    "contato": "Maria",
    "telefone": "(11) 3333-4444",
    "email": "vendas@alfa.com",
    "numero": "OR-555",
    "data": "2026-09-15",
    "pagamento": "28 dias",
    "entrega": "5 dias úteis",
    "frete": "cif",
    "valorFrete": "1.234,50",
    "obs": "",
    "itens": [
        {"produto": "Dipirona 500mg", "marca": "EMS", "unidade": "cx", "qtd": 10, "valorUnit": 12.34567},
        {"produto": "", "marca": "", "unidade": "", "qtd": 0, "valorUnit": 0},
    ],
}


class ClienteFalso:
    def __init__(self, resposta=None, erro=None):
        self.chamadas = []
        self.resposta = resposta if resposta is not None else RESPOSTA_IA
        self.erro = erro
        self.messages = self

    def create(self, **kwargs):
        self.chamadas.append(kwargs)
        if self.erro:
            raise self.erro
        return SimpleNamespace(
            content=[SimpleNamespace(type="tool_use", name="registrar_cotacao", input=self.resposta)],
            usage=SimpleNamespace(input_tokens=1500, output_tokens=300),
            stop_reason="tool_use",
        )


@pytest.fixture
def ia_falsa(monkeypatch):
    falso = ClienteFalso()
    monkeypatch.setattr(leitor, "obter_cliente", lambda: falso)
    return falso


def _blocos(falso):
    return falso.chamadas[0]["messages"][0]["content"]


def test_texto_colado(cliente, ia_falsa, db):
    r = cliente.post("/api/ia/ler-cotacao", data={"texto": "Dipirona cx 12,34"})
    assert r.status_code == 200, r.text
    d = r.json()
    assert d["fornecedor"] == "Distribuidora Alfa"
    assert d["frete"] == "CIF"
    assert d["valorFrete"] == 1234.5
    assert d["itens"] == [{"produto": "Dipirona 500mg", "marca": "EMS", "unidade": "CX", "qtd": 10.0, "valorUnit": 12.3457}]

    chamada = ia_falsa.chamadas[0]
    assert chamada["model"] == "claude-sonnet-5-5"
    assert chamada["tool_choice"] == {"type": "auto"}
    assert chamada["tools"][0]["strict"] is True
    assert "26.643.172/0001-77" in chamada["system"]
    assert "Dipirona cx 12,34" in _blocos(ia_falsa)[0]["text"]

    uso = db.query(UsoIA).one()
    assert (uso.tokens_entrada, uso.tokens_saida, uso.tipo_entrada, uso.sucesso) == (1500, 300, "texto", True)
    # Nada é salvo como registro.
    assert cliente.get("/api/registros").json()["total"] == 0


def test_pdf_e_imagem_vao_como_bloco(cliente, ia_falsa):
    r = cliente.post("/api/ia/ler-cotacao", files={"arquivo": ("cot.pdf", b"%PDF-1.4 teste", "application/pdf")})
    assert r.status_code == 200
    assert _blocos(ia_falsa)[0]["type"] == "document"
    assert _blocos(ia_falsa)[0]["source"]["media_type"] == "application/pdf"

    r = cliente.post("/api/ia/ler-cotacao", files={"arquivo": ("foto.JPG", b"\xff\xd8\xff imagem", "image/jpeg")})
    assert r.status_code == 200
    bloco = ia_falsa.chamadas[1]["messages"][0]["content"][0]
    assert bloco["type"] == "image" and bloco["source"]["media_type"] == "image/jpeg"


def test_docx_inclui_tabelas(cliente, ia_falsa):
    doc = Document()
    doc.add_paragraph("Orçamento 99 - Loja Gama")
    t = doc.add_table(rows=2, cols=3)
    for c, v in zip(t.rows[0].cells, ("Produto", "Qtd", "Valor")):
        c.text = v
    for c, v in zip(t.rows[1].cells, ("Luva M", "100", "0,45")):
        c.text = v
    buf = io.BytesIO()
    doc.save(buf)
    r = cliente.post("/api/ia/ler-cotacao", files={"arquivo": ("cot.docx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200
    texto = _blocos(ia_falsa)[0]["text"]
    assert "Orçamento 99 - Loja Gama" in texto
    assert "Luva M | 100 | 0,45" in texto


def test_xlsx_cada_aba_em_csv(cliente, ia_falsa):
    wb = Workbook()
    wb.active.title = "Cotação"
    wb.active.append(["Produto", "Preço"])
    wb.active.append(["Soro 500ml", 4.5])
    wb.create_sheet("Frete").append(["CIF", 0])
    buf = io.BytesIO()
    wb.save(buf)
    r = cliente.post("/api/ia/ler-cotacao", files={"arquivo": ("cot.xlsx", buf.getvalue(), "application/octet-stream")})
    assert r.status_code == 200
    texto = _blocos(ia_falsa)[0]["text"]
    assert "### Aba: Cotação" in texto and "Soro 500ml,4.5" in texto and "### Aba: Frete" in texto


def test_csv(cliente, ia_falsa):
    r = cliente.post("/api/ia/ler-cotacao", files={"arquivo": ("c.csv", "produto;valor\nGaze;1,20".encode("cp1252"), "text/csv")})
    assert r.status_code == 200
    assert "Gaze;1,20" in _blocos(ia_falsa)[0]["text"]


def test_rejeita_doc_antigo(cliente, ia_falsa):
    r = cliente.post("/api/ia/ler-cotacao", files={"arquivo": ("velho.doc", b"xx", "application/msword")})
    assert r.status_code == 422
    assert ".docx ou PDF" in r.json()["detail"]
    assert ia_falsa.chamadas == []


def test_rejeita_tipo_e_tamanho(cliente, ia_falsa, monkeypatch):
    assert cliente.post("/api/ia/ler-cotacao", files={"arquivo": ("a.zip", b"x", "application/zip")}).status_code == 422
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "ia_max_mb", 1)
    r = cliente.post("/api/ia/ler-cotacao", files={"arquivo": ("g.pdf", b"%PDF" + b"0" * (1024 * 1024), "application/pdf")})
    assert r.status_code == 422
    assert "1 MB" in r.json()["detail"]
    assert ia_falsa.chamadas == []


def test_sem_conteudo(cliente, ia_falsa):
    r = cliente.post("/api/ia/ler-cotacao", data={"texto": "   "})
    assert r.status_code == 422
    assert "arquivo ou cole o texto" in r.json()["detail"]


def test_nunca_usa_h3_como_fornecedor(cliente, monkeypatch):
    falso = ClienteFalso(resposta={**RESPOSTA_IA, "fornecedor": "H3 Pharma Comércio e Serviços Ltda", "cnpj": "26.643.172/0001-77", "data": "15/09/2026"})
    monkeypatch.setattr(leitor, "obter_cliente", lambda: falso)
    d = cliente.post("/api/ia/ler-cotacao", data={"texto": "x"}).json()
    assert d["fornecedor"] == "" and d["cnpj"] == "" and d["data"] == ""


def test_erro_da_api_em_portugues(cliente, monkeypatch, db):
    import anthropic
    import httpx

    req = httpx.Request("POST", "https://api.anthropic.com/v1/messages")
    erro = anthropic.RateLimitError("limite", response=httpx.Response(429, request=req), body=None)
    monkeypatch.setattr(leitor, "obter_cliente", lambda: ClienteFalso(erro=erro))
    r = cliente.post("/api/ia/ler-cotacao", data={"texto": "x"})
    assert r.status_code == 429
    assert "Aguarde" in r.json()["detail"]
    assert db.query(UsoIA).one().sucesso is False


def test_sem_chave_configurada(cliente, monkeypatch):
    from app.config import get_settings

    monkeypatch.setattr(get_settings(), "anthropic_api_key", "")
    r = cliente.post("/api/ia/ler-cotacao", data={"texto": "x"})
    assert r.status_code == 503
    assert "não está configurada" in r.json()["detail"]


def test_resumo_uso_ia_admin(admin, cliente, ia_falsa):
    cliente.post("/api/ia/ler-cotacao", data={"texto": "a"})
    admin.post("/api/ia/ler-cotacao", data={"texto": "b"})
    d = admin.get("/api/admin/uso-ia").json()
    assert d["total"]["leituras"] == 2
    assert d["total"]["tokens_entrada"] == 3000
    assert d["total"]["custo_estimado_usd"] == round(3000 / 1e6 * 2 + 600 / 1e6 * 10, 4)
    assert {u["nome"] for u in d["por_unidade"]} == {"Matriz", "Filial Norte"}
