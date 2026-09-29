"""Converte o arquivo enviado em blocos de conteúdo para a API da Anthropic."""

import base64
import io
from dataclasses import dataclass


class ArquivoInvalido(Exception):
    """Erro com mensagem pronta para mostrar ao usuário."""


LIMITE_TEXTO = 300_000  # caracteres enviados à IA (planilhas muito grandes são cortadas)

IMAGENS = {
    ".jpg": "image/jpeg",
    ".jpeg": "image/jpeg",
    ".png": "image/png",
    ".webp": "image/webp",
}
EXTENSOES_ACEITAS = (".pdf", *IMAGENS, ".docx", ".xlsx", ".xls", ".csv")


@dataclass
class Conteudo:
    tipo: str  # pdf | imagem | word | planilha | texto
    blocos: list[dict]


def _extensao(nome: str) -> str:
    nome = (nome or "").lower().strip()
    return nome[nome.rfind(".") :] if "." in nome else ""


def _bloco_texto(titulo: str, texto: str) -> dict:
    texto = texto.strip()
    if not texto:
        raise ArquivoInvalido("Não encontrei texto nesse arquivo. Confira se ele não está vazio.")
    if len(texto) > LIMITE_TEXTO:
        texto = texto[:LIMITE_TEXTO] + "\n[... conteúdo cortado por ser muito grande ...]"
    return {"type": "text", "text": f"{titulo}:\n\n{texto}"}


def texto_docx(dados: bytes) -> str:
    from docx import Document
    from docx.table import Table
    from docx.text.paragraph import Paragraph

    doc = Document(io.BytesIO(dados))
    linhas: list[str] = []

    def tabela(t: Table) -> None:
        for row in t.rows:
            celulas: list[str] = []
            anterior = None
            for cell in row.cells:
                # Células mescladas se repetem em python-docx; evita duplicar.
                if cell._tc is anterior:
                    continue
                anterior = cell._tc
                celulas.append(" ".join(cell.text.split()))
            if any(celulas):
                linhas.append(" | ".join(celulas))
        linhas.append("")

    # Percorre o corpo na ordem original (parágrafos e tabelas intercalados).
    for filho in doc.element.body.iterchildren():
        tag = filho.tag.rsplit("}", 1)[-1]
        if tag == "p":
            texto = Paragraph(filho, doc).text.strip()
            if texto:
                linhas.append(texto)
        elif tag == "tbl":
            tabela(Table(filho, doc))
    for secao in doc.sections:
        for parte in (secao.header, secao.footer):
            for p in parte.paragraphs:
                if p.text.strip():
                    linhas.append(p.text.strip())
            for t in parte.tables:
                tabela(t)
    return "\n".join(linhas)


def _decodificar(dados: bytes) -> str:
    for codificacao in ("utf-8-sig", "cp1252", "latin-1"):
        try:
            return dados.decode(codificacao)
        except UnicodeDecodeError:
            continue
    return dados.decode("utf-8", errors="replace")


def texto_planilha(dados: bytes, extensao: str) -> str:
    import pandas as pd

    if extensao == ".csv":
        return _decodificar(dados)
    motor = "xlrd" if extensao == ".xls" else "openpyxl"
    abas = pd.read_excel(io.BytesIO(dados), sheet_name=None, header=None, dtype=str, engine=motor)
    partes = []
    for nome, df in abas.items():
        df = df.dropna(how="all").dropna(axis=1, how="all")
        if df.empty:
            continue
        partes.append(f"### Aba: {nome}\n{df.to_csv(index=False, header=False)}")
    return "\n\n".join(partes)


def preparar_arquivo(nome: str, dados: bytes, limite_mb: int) -> Conteudo:
    ext = _extensao(nome)
    if ext == ".doc":
        raise ArquivoInvalido(
            "Arquivos .doc (Word antigo) não são aceitos. Abra no Word e salve como .docx ou PDF."
        )
    if ext not in EXTENSOES_ACEITAS:
        raise ArquivoInvalido(
            "Tipo de arquivo não aceito. Envie PDF, imagem (JPG, PNG, WebP), Word (.docx) "
            "ou planilha (.xlsx, .xls, .csv)."
        )
    if len(dados) > limite_mb * 1024 * 1024:
        raise ArquivoInvalido(f"O arquivo passa do limite de {limite_mb} MB.")
    if not dados:
        raise ArquivoInvalido("O arquivo está vazio.")

    if ext == ".pdf":
        if not dados.startswith(b"%PDF"):
            raise ArquivoInvalido("Esse arquivo não parece ser um PDF válido.")
        b64 = base64.standard_b64encode(dados).decode()
        return Conteudo(
            "pdf", [{"type": "document", "source": {"type": "base64", "media_type": "application/pdf", "data": b64}}]
        )
    if ext in IMAGENS:
        b64 = base64.standard_b64encode(dados).decode()
        return Conteudo(
            "imagem", [{"type": "image", "source": {"type": "base64", "media_type": IMAGENS[ext], "data": b64}}]
        )
    try:
        if ext == ".docx":
            return Conteudo("word", [_bloco_texto("Conteúdo do documento Word", texto_docx(dados))])
        return Conteudo("planilha", [_bloco_texto("Conteúdo da planilha (CSV)", texto_planilha(dados, ext))])
    except ArquivoInvalido:
        raise
    except Exception as exc:  # arquivo corrompido ou com formato diferente da extensão
        raise ArquivoInvalido(
            "Não consegui abrir esse arquivo. Confira se ele não está corrompido ou protegido por senha."
        ) from exc


def preparar_texto(texto: str) -> Conteudo:
    return Conteudo("texto", [_bloco_texto("Texto da cotação colado pelo usuário", texto)])
